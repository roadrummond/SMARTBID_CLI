#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SmartBid CLI v1.0 — Plataforma de Inteligência em Licitações
R2 Tecnologia Ltda.

Uso:  python smartbid.py
Requisitos: pip install -r requirements.txt  |  variável ANTHROPIC_API_KEY definida
"""
import base64
import csv
import datetime as dt
import json
import os
import re
import sys
import time
import zipfile
from pathlib import Path

BASE = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
CFG_FILE = BASE / "config.json"
OUT = BASE / "saida"
OUT.mkdir(exist_ok=True)
INTEL_CSV = OUT / "smartintel.csv"
HIST_SEARCH = OUT / "historico_search.json"

DEFAULT_CFG = {
    "modelo": "claude-sonnet-5-5",
    "max_tokens": 32000,
    "produto": "ExaGrid Tiered Backup Storage",
    "pasta_produto": "docs_produto",
    "keywords": ["appliance de backup", "solução de backup", "proteção contra ransomware",
                 "tiered backup", "storage de backup", "backup imutável"],
    "ufs": [],
    "valor_min": 0,
    "limite_caracteres": 450000,
}

DISCLAIMER = ("AVISO LEGAL (SmartLegal): os textos produzidos pelo SmartBid têm caráter informativo e de "
              "apoio técnico, não constituindo assessoria jurídica. O protocolo final de qualquer documento é "
              "responsabilidade do usuário ou de seu advogado habilitado.")

SESSAO = {"cliente": None, "pregao": None, "contexto": {}}


# ───────────────────────────── utilidades ─────────────────────────────
def cfg():
    if not CFG_FILE.exists():
        CFG_FILE.write_text(json.dumps(DEFAULT_CFG, indent=2, ensure_ascii=False), encoding="utf-8")
    c = dict(DEFAULT_CFG)
    c.update(json.loads(CFG_FILE.read_text(encoding="utf-8")))
    return c


def linha(ch="─", n=70):
    print(ch * n)


def perguntar(msg, padrao=None, obrigatorio=False):
    while True:
        sfx = f" [{padrao}]" if padrao not in (None, "") else ""
        v = input(f"{msg}{sfx}: ").strip()
        if not v and padrao is not None:
            return str(padrao)
        if v or not obrigatorio:
            return v
        print("  ! Campo obrigatório.")


def num(msg, padrao):
    while True:
        v = perguntar(msg, padrao).replace(",", ".")
        try:
            return float(v)
        except ValueError:
            print("  ! Informe um número.")


def caminho(msg, obrigatorio=True):
    while True:
        v = perguntar(msg + " (arraste o arquivo aqui ou Enter p/ pular)" if not obrigatorio else msg,
                      obrigatorio=obrigatorio)
        if not v:
            return None
        p = Path(v.strip().strip('"').strip("'"))
        if p.exists():
            return p
        print(f"  ! Arquivo não encontrado: {p}")


def slug(s):
    s = re.sub(r"[^\w]+", "", (s or "").upper().replace(" ", "_").replace("/", "")) or "NA"
    return s


def nome_saida(assunto, ext="docx"):
    return OUT / f"{slug(SESSAO['cliente'])}_{slug(SESSAO['pregao'])}_{assunto}_{dt.date.today():%d%m%y}.{ext}"


def dados_sessao():
    if not SESSAO["cliente"]:
        SESSAO["cliente"] = perguntar("Cliente / órgão (sigla, ex.: MPPT)", obrigatorio=True)
    if not SESSAO["pregao"]:
        SESSAO["pregao"] = perguntar("Nº do pregão (ex.: PE 13/2026)", obrigatorio=True)


def dias_uteis(data, n):
    """Soma (n>0) ou subtrai (n<0) dias úteis — não considera feriados."""
    passo = 1 if n > 0 else -1
    d, cont = data, 0
    while cont < abs(n):
        d += dt.timedelta(days=passo)
        if d.weekday() < 5:
            cont += 1
    return d


def ler_data(msg):
    while True:
        v = perguntar(msg + " (dd/mm/aaaa)", obrigatorio=True)
        try:
            return dt.datetime.strptime(v, "%d/%m/%Y").date()
        except ValueError:
            print("  ! Data inválida.")


# ───────────────────────────── leitura de arquivos ─────────────────────────────
def extrair(p: Path):
    """Retorna (texto, pdf_bytes_se_escaneado)."""
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(str(p))
        partes = []
        for i, pg in enumerate(r.pages, 1):
            try:
                t = pg.extract_text() or ""
            except Exception:
                t = ""
            partes.append(f"\n[Página {i}]\n{t}")
        txt = "".join(partes)
        escaneado = len(re.sub(r"\s|\[Página \d+\]", "", txt)) < 150 * max(1, len(r.pages)) * 0.2
        return txt, (p.read_bytes() if escaneado else None)
    if ext == ".docx":
        from docx import Document
        d = Document(str(p))
        partes = [par.text for par in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                partes.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(partes), None
    if ext == ".zip":
        txts = []
        with zipfile.ZipFile(p) as z:
            destino = p.with_suffix("")
            z.extractall(destino)
        for f in sorted(destino.rglob("*")):
            if f.suffix.lower() in (".pdf", ".docx", ".txt", ".md", ".csv"):
                t, _ = extrair(f)
                txts.append(f"\n===== {f.name} =====\n{t}")
        return "".join(txts), None
    return p.read_text(encoding="utf-8", errors="ignore"), None


def bloco(rotulo, p: Path, limite):
    print(f"  · lendo {rotulo}: {p.name}")
    txt, pdf = extrair(p)
    if pdf is not None and len(pdf) < 30 * 1024 * 1024:
        print("    (PDF sem camada de texto — enviando como documento para OCR)")
        return [{"type": "text", "text": f"===== {rotulo}: {p.name} ====="},
                {"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                "data": base64.b64encode(pdf).decode()}}]
    if len(txt) > limite:
        print(f"    (texto truncado para {limite:,} caracteres)")
        txt = txt[:limite]
    return [{"type": "text", "text": f"===== {rotulo}: {p.name} =====\n{txt}"}]


def docs_produto(c, limite):
    pasta = BASE / c["pasta_produto"]
    if not pasta.exists():
        pasta.mkdir(exist_ok=True)
    arqs = [f for f in sorted(pasta.iterdir()) if f.suffix.lower() in (".pdf", ".docx", ".txt", ".md")]
    if not arqs:
        print(f"  ! Nenhuma documentação de produto em {pasta} (coloque datasheets/manuais oficiais lá).")
        return []
    blocos = []
    por = max(20000, limite // max(1, len(arqs)))
    for f in arqs:
        blocos += bloco("DOCUMENTAÇÃO OFICIAL DO PRODUTO", f, por)
    return blocos


# ───────────────────────────── Claude ─────────────────────────────
SMARTLEGAL = """Você é o SmartBid, plataforma de inteligência em licitações públicas brasileiras, com o filtro
normativo SmartLegal sempre ativo. Regras obrigatórias:
1. Fundamente todo argumento em dispositivo específico da Lei 14.133/2021 (art. 5º princípios; art. 9º, I, "a"
   vedação a restrição de competitividade; art. 41 e art. 74 §1º indicação de marca; art. 55 prazos e §1º reabertura
   após alteração; art. 67 qualificação técnica; art. 164 impugnação/esclarecimento até 3 dias úteis antes da
   abertura; art. 165 recurso em 3 dias úteis, reconsideração em 3 dias úteis e decisão superior em 10 dias úteis;
   contrarrazões no mesmo prazo do recurso a partir da intimação). Obrigatoriedade de publicação no PNCP (art. 54).
2. Acórdãos TCU de referência fornecidos pelo usuário: 1.604/2025 (limite de 50% em atestados), 1.712/2025 e
   1.091/2025 (normas/certificações só com essencialidade motivada). Não invente acórdãos, números ou ementas: se
   não tiver certeza, escreva "[conferir jurisprudência]".
3. Conteúdo sem base normativa deve ser marcado "[SEM FUNDAMENTO LEGAL IDENTIFICADO]".
4. Use apenas informações presentes nos documentos fornecidos. Nunca invente especificações, preços, páginas ou
   links. Comprovação técnica: somente documentação OFICIAL do fabricante (nunca revenda/parceiro).
5. Cite a página do edital ([Página N]) sempre que possível.
6. Responda em português do Brasil, em Markdown (títulos #, tabelas |, listas -). Seja objetivo e completo."""


def chamar_claude(instrucao, blocos, c):
    try:
        import anthropic
    except ImportError:
        sys.exit("Instale as dependências: pip install -r requirements.txt")
    if not os.environ.get("ANTHROPIC_API_KEY"):
        sys.exit("Defina a variável de ambiente ANTHROPIC_API_KEY.")
    cli = anthropic.Anthropic()
    conteudo = blocos + [{"type": "text", "text": instrucao}]
    print("  · processando ", end="", flush=True)
    t0, n = time.time(), 0
    with cli.messages.stream(model=c["modelo"], max_tokens=int(c["max_tokens"]), system=SMARTLEGAL,
                             messages=[{"role": "user", "content": conteudo}]) as s:
        for _ in s.text_stream:
            n += 1
            if n % 40 == 0:
                print("▮", end="", flush=True)
        msg = s.get_final_message()
    print(f"  ok ({time.time() - t0:.0f}s)")
    return "".join(b.text for b in msg.content if b.type == "text")


# ───────────────────────────── relatório DOCX ─────────────────────────────
def _runs(par, texto):
    for i, parte in enumerate(re.split(r"\*\*", texto)):
        if parte:
            r = par.add_run(parte)
            r.bold = i % 2 == 1


def salvar(md, assunto, titulo):
    from docx import Document
    from docx.shared import Pt
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading(titulo, 0)
    d.add_paragraph(f"Cliente: {SESSAO['cliente']}   |   Pregão: {SESSAO['pregao']}   |   "
                    f"Emissão: {dt.datetime.now():%d/%m/%Y %H:%M}")
    ls = md.splitlines()
    i = 0
    while i < len(ls):
        l = ls[i].rstrip()
        if l.strip().startswith("|"):
            tab = []
            while i < len(ls) and ls[i].strip().startswith("|"):
                cells = [x.strip() for x in ls[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", x or "--") for x in cells):
                    tab.append(cells)
                i += 1
            if tab:
                ncol = max(len(r) for r in tab)
                t = d.add_table(rows=0, cols=ncol)
                t.style = "Table Grid"
                for ri, r in enumerate(tab):
                    row = t.add_row().cells
                    for ci in range(ncol):
                        p = row[ci].paragraphs[0]
                        _runs(p, r[ci] if ci < len(r) else "")
                        if ri == 0:
                            for run in p.runs:
                                run.bold = True
                d.add_paragraph()
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", l)
        if m:
            d.add_heading(m.group(2).replace("**", ""), min(len(m.group(1)), 4))
        elif re.match(r"^\s*[-*]\s+", l):
            _runs(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", l))
        elif re.match(r"^\s*\d+[.)]\s+", l):
            _runs(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", l))
        elif l.strip() in ("---", "***"):
            pass
        elif l.strip():
            _runs(d.add_paragraph(), l)
        i += 1
    p = d.add_paragraph()
    r = p.add_run(DISCLAIMER)
    r.italic = True
    r.font.size = Pt(8)
    arq = nome_saida(assunto)
    d.save(arq)
    arq.with_suffix(".md").write_text(md + "\n\n" + DISCLAIMER, encoding="utf-8")
    print(f"\n  ✔ Relatório gerado: {arq}")
    return arq


def mostrar(md, n=40):
    linha()
    for l in md.splitlines()[:n]:
        print(l)
    if len(md.splitlines()) > n:
        print(f"... ({len(md.splitlines()) - n} linhas adicionais no relatório)")
    linha()


# ───────────────────────────── prompts dos módulos ─────────────────────────────
P_INSIGHT = """MÓDULO SmartInsight — Síntese executiva do edital. NÃO interprete nem avalie; apenas organize.
Estrutura obrigatória:
# 1. Resumo do certame — tabela (Campo | Informação): Órgão, UASG, Nº do pregão/processo, Objeto, Modalidade,
  Critério de julgamento, Modo de disputa, Valor estimado (por item/lote se houver), Data/hora da abertura,
  Portal, Link do PNCP (se constar no edital; senão "não informado no edital").
# 2. Prazos críticos — tabela (Evento | Data/prazo | Fundamento/página). Inclua a data-limite para impugnação e
  esclarecimento (3 dias úteis antes da abertura — art. 164), prazo de entrega, instalação, vigência, garantia,
  visita técnica.
# 3. Exigências técnicas de HARDWARE (lista com página)
# 4. Exigências técnicas de SOFTWARE / licenciamento / serviços (lista com página)
# 5. Documentos de habilitação exigidos (jurídica, fiscal, econômica, técnica — destaque atestados e percentuais)
# 6. Penalidades, multas, SLA e cláusulas de risco (tabela)
# 7. Resumo de uma página para diretoria (máx. 15 linhas)."""

P_ANALYSE = """MÓDULO SmartAnalyse — Análise estratégica e score de vencibilidade para o produto: {produto}.
Use a documentação oficial do produto fornecida para avaliar aderência.
# 1. Resumo (tabela: Órgão, Pregão, Objeto, Valor, Abertura, Link PNCP)
# 2. Aderência técnica — tabela (Requisito resumido | Página | Aderência: Plena/Parcial/Não atende | Observação)
# 3. Riscos — três tabelas: Técnicos, Jurídicos, Administrativos (Risco | Página | Impacto Alto/Médio/Baixo | Mitigação)
# 4. Direcionamento técnico — cláusulas que limitam a competitividade (art. 9º, I, "a"), indicação de marca
  (art. 41), certificações sem motivação, atestados acima de 50%; indique quantos fabricantes cada cláusula
  tende a eliminar SOMENTE se for verificável, sem nomeá-los.
# 5. Indícios de redação gerada por IA — padrões observados e grau de confiança (indicativo, não conclusivo).
# 6. Score de vencibilidade 0–10 — tabela por categoria (Técnica 40%, Jurídica 20%, Administrativa 15%,
  Competitiva 25%) com nota e justificativa; nota final ponderada.
# 7. Recomendação: PARTICIPAR / PARTICIPAR COM RESSALVAS / NÃO PARTICIPAR — com as ações prévias necessárias e a
  data-limite para questionamentos (art. 164)."""

P_COMPLIANCE = """MÓDULO SmartCompliance — Matriz de conformidade técnica para o produto: {produto}.
# 1. Resumo (tabela: Órgão, Pregão, Objeto, Link PNCP, Modelo recomendado que atende o edital)
# 2. Matriz de conformidade — UMA tabela com exatamente estas colunas:
  | Item | Requisito (transcrição literal do edital) | Atende | Como atende | Comprovação |
  - Item: numeração do edital. Requisito: copie o texto exato, sem resumir.
  - Atende: "Atende Plenamente", "Atende Parcialmente" ou "Não Atende".
  - Como atende: frase objetiva e positiva, mostrando como o requisito é resolvido, com o resultado para o órgão
    como sujeito (ex.: "as operações de backup continuam nos nodes remanescentes"). Não exponha deficiências.
  - Comprovação: documento OFICIAL do fabricante fornecido + página (e URL oficial se constar no documento).
    Se não houver comprovação nos documentos: "Comprovação pendente — localizar em documentação oficial".
  Cubra TODOS os requisitos técnicos do objeto (hardware e software).
# 3. Questionamentos (pedidos de esclarecimento — art. 164) SOMENTE para itens Parcial/Não Atende.
  NÃO questione item que já é atendido como está escrito. Cada questionamento, curto e objetivo, na ordem:
  (1) Requisito do edital (item e trecho); (2) Embasamento técnico (2–4 linhas, persuasivo, neutro, sem citar
  fabricantes, com fundamento legal); (3) Pergunta redigida para ser respondida "SIM", na forma "Entendemos
  que ... está correto o nosso entendimento?".
# 4. Impugnação — somente se houver direcionamento claro: minuta enxuta (endereçamento, tempestividade, fatos,
  fundamentos, pedido de nova redação e reabertura de prazo — art. 55 §1º). Se não houver, escreva "Não aplicável".
# 5. Data-limite para protocolo (3 dias úteis antes da abertura)."""

P_POSITION = """MÓDULO SmartPosition — Inteligência competitiva e posicionamento para o produto: {produto}.
Base histórica: use EXCLUSIVAMENTE o histórico SmartIntel fornecido (se houver). Nunca invente preços,
vencedores ou margens; sem histórico, declare "sem dados históricos" e trabalhe apenas com o TR.
# 1. Histórico de pregões similares (tabela: Órgão | Pregão | Vencedor | Produto | Preço | Resultado)
# 2. Faixa de preço e margem praticadas (somente com dados do histórico; mostre o cálculo)
# 3. Perfil de arquiteturas concorrentes que o TR favorece ou desfavorece (sem nomear fabricantes no texto que
  irá ao órgão; na análise interna pode nomear tipos de solução)
# 4. Pontos fortes da solução proposta frente a este TR (tabela Requisito | Vantagem | Benefício para o órgão)
# 5. Questionamentos sugeridos para melhorar o cenário competitivo (estrutura curta: requisito, embasamento,
  pergunta)
# 6. Recomendação de posicionamento técnico e comercial (bullet points)."""

P_PROPOSAL = """MÓDULO SmartProposal — Proposta técnica e comercial para o produto: {produto}.
Dados comerciais informados pelo usuário:
{comercial}
Use a matriz de conformidade (se fornecida) como base. Não invente valores: campos sem dado = "[PREENCHER]".
# 1. Carta de apresentação da proposta (modelo do edital, se houver)
# 2. Proposta técnica — descrição da solução, item a item conforme o edital, com modelo, quantidades e garantia
# 3. Declarações obrigatórias exigidas pelo edital (liste cada uma e redija o texto com campos [PREENCHER])
# 4. Proposta comercial — tabela (Item | Descrição | Qtd | Valor unitário | Valor total) e memória de cálculo
# 5. Diferenciais técnicos (benefício para o órgão como sujeito)
# 6. Checklist de consistência proposta × edital (Requisito | Onde está atendido na proposta | OK/Pendente)."""

P_REVIEW = """MÓDULO SmartReview — Análise da proposta do concorrente vencedor e recurso administrativo (art. 165).
Data-limite calculada para interposição: {prazo}.
# 1. Resumo (Órgão, Pregão, Concorrente, Produto ofertado, Valor)
# 2. Análise item a item — tabela (Item | Requisito do TR | O que o concorrente ofertou/comprovou | Conformidade |
  Evidência: página da proposta/catálogo)
# 3. Não conformidades — ordenadas por relevância/impacto (técnicas, documentais, jurídicas), cada uma com
  fundamento (vinculação ao edital e julgamento objetivo — art. 5º; demais artigos aplicáveis)
# 4. Minuta do RECURSO ADMINISTRATIVO — endereçamento, qualificação [PREENCHER], tempestividade, síntese dos fatos,
  razões (uma seção por não conformidade, da mais forte para a mais fraca), pedidos (desclassificação/inabilitação
  e convocação do próximo colocado), local/data/assinatura [PREENCHER].
Somente aponte não conformidade comprovável pelos documentos fornecidos."""

P_DEFENSE = """MÓDULO SmartDefense — Contrarrazões ao recurso do concorrente (art. 165).
Data-limite calculada para contrarrazões: {prazo}.
# 1. Estrutura do recurso — tabela (Nº | Alegação do recorrente | Tipo: técnica/jurídica/documental | Página)
# 2. Contra-argumentação por ponto — tabela (Alegação | Resposta | Prova/página | Fundamento legal | Risco de
  provimento Baixo/Médio/Alto)
# 3. Risco global de provimento (baixo/médio/alto) com justificativa
# 4. Reforço documental recomendado (o que anexar)
# 5. Minuta das CONTRARRAZÕES — endereçamento, qualificação [PREENCHER], tempestividade, síntese do recurso,
  razões ponto a ponto, pedido de não provimento e manutenção do resultado, local/data/assinatura [PREENCHER]."""


# ───────────────────────────── execução dos módulos ─────────────────────────────
def rodar(assunto, titulo, instrucao, blocos, c, chave=None):
    md = chamar_claude(instrucao, blocos, c)
    if chave:
        SESSAO["contexto"][chave] = md
    mostrar(md)
    return salvar(md, assunto, titulo), md


def pedir_edital(c, edital=None):
    edital = edital or caminho("Arquivo do EDITAL/TR (PDF, DOCX, ZIP)")
    return edital, bloco("EDITAL / TERMO DE REFERÊNCIA", edital, c["limite_caracteres"])


def ctx(chave, rotulo):
    md = SESSAO["contexto"].get(chave)
    return [{"type": "text", "text": f"===== {rotulo} (gerado anteriormente) =====\n{md}"}] if md else []


def m_insight(c, edital=None):
    dados_sessao()
    _, b = pedir_edital(c, edital)
    return rodar("SINTESE_EXECUTIVA", "SmartInsight — Síntese Executiva do Edital", P_INSIGHT, b, c, "insight")


def m_analyse(c, edital=None):
    dados_sessao()
    _, b = pedir_edital(c, edital)
    b += docs_produto(c, 150000) + ctx("insight", "SÍNTESE EXECUTIVA")
    return rodar("ANALISE_VENCIBILIDADE", "SmartAnalyse — Análise Estratégica e Score de Vencibilidade",
                 P_ANALYSE.format(produto=c["produto"]), b, c, "analyse")


def m_compliance(c, edital=None):
    dados_sessao()
    _, b = pedir_edital(c, edital)
    b += docs_produto(c, 200000) + ctx("analyse", "ANÁLISE SMARTANALYSE")
    return rodar("MATRIZ_CONF_TECNICA", "SmartCompliance — Matriz de Conformidade e Questionamentos",
                 P_COMPLIANCE.format(produto=c["produto"]), b, c, "compliance")


def intel_texto():
    if not INTEL_CSV.exists():
        return []
    return [{"type": "text", "text": "===== HISTÓRICO SMARTINTEL =====\n" +
             INTEL_CSV.read_text(encoding="utf-8")}]


def m_position(c, edital=None):
    dados_sessao()
    _, b = pedir_edital(c, edital)
    extra = caminho("Histórico adicional de pregões (CSV/PDF/atas)", obrigatorio=False)
    if extra:
        b += bloco("HISTÓRICO ADICIONAL", extra, 100000)
    b += intel_texto() + docs_produto(c, 80000) + ctx("compliance", "MATRIZ DE CONFORMIDADE")
    return rodar("POSICIONAMENTO_COMPETITIVO", "SmartPosition — Inteligência Competitiva",
                 P_POSITION.format(produto=c["produto"]), b, c, "position")


def m_proposal(c, edital=None):
    dados_sessao()
    _, b = pedir_edital(c, edital)
    if "compliance" not in SESSAO["contexto"]:
        mat = caminho("Matriz de conformidade já gerada (DOCX/MD)", obrigatorio=False)
        if mat:
            b += bloco("MATRIZ DE CONFORMIDADE", mat, 150000)
    b += ctx("compliance", "MATRIZ DE CONFORMIDADE") + docs_produto(c, 100000)
    print("  Dados comerciais (Enter = [PREENCHER]):")
    com = {k: perguntar(f"    {k}") or "[PREENCHER]" for k in
           ("Razão social / CNPJ do licitante", "Modelo ofertado", "Valor unitário por item (R$)",
            "Validade da proposta (dias)", "Prazo de entrega", "Garantia (meses)")}
    comercial = "\n".join(f"- {k}: {v}" for k, v in com.items())
    return rodar("PROPOSTA_TEC_COMERCIAL", "SmartProposal — Proposta Técnica e Comercial",
                 P_PROPOSAL.format(produto=c["produto"], comercial=comercial), b, c, "proposal")


def m_review(c, edital=None):
    dados_sessao()
    _, b = pedir_edital(c, edital)
    prop = caminho("PROPOSTA/catálogos do CONCORRENTE vencedor (PDF/ZIP)")
    b += bloco("PROPOSTA DO CONCORRENTE", prop, 250000)
    ata = caminho("Ata/decisão do pregão", obrigatorio=False)
    if ata:
        b += bloco("ATA / DECISÃO", ata, 60000)
    d = ler_data("Data da intimação / lavratura da ata")
    prazo = f"{dias_uteis(d, 3):%d/%m/%Y} (3 dias úteis — art. 165; feriados não considerados, conferir)"
    print(f"  ⏰ Prazo para recurso: {prazo}")
    return rodar("RECURSO_ADMINISTRATIVO", "SmartReview — Análise do Concorrente e Recurso",
                 P_REVIEW.format(prazo=prazo), b, c, "review")


def m_defense(c, edital=None):
    dados_sessao()
    _, b = pedir_edital(c, edital)
    rec = caminho("RECURSO interposto pelo concorrente (PDF/DOCX)")
    b += bloco("RECURSO DO CONCORRENTE", rec, 150000)
    prop = caminho("Nossa proposta (PDF/DOCX)", obrigatorio=False)
    if prop:
        b += bloco("NOSSA PROPOSTA", prop, 150000)
    b += docs_produto(c, 100000)
    d = ler_data("Data da intimação da interposição do recurso")
    prazo = f"{dias_uteis(d, 3):%d/%m/%Y} (3 dias úteis; feriados não considerados, conferir)"
    print(f"  ⏰ Prazo para contrarrazões: {prazo}")
    return rodar("CONTRARRAZOES", "SmartDefense — Contrarrazões ao Recurso",
                 P_DEFENSE.format(prazo=prazo), b, c, "defense")


# ───────────────────────────── SmartTools (cálculo local) ─────────────────────────────
def m_sizing(c):
    dados_sessao()
    print("  SmartSizing — premissas (TB decimais):")
    fe = num("    Dados primários protegidos hoje (TB)", 100)
    cresc = num("    Crescimento anual (%)", 20) / 100
    anos = int(num("    Horizonte (anos)", 5))
    chg = num("    Taxa de mudança diária (%)", 2) / 100
    dd = int(num("    Retenção diária (dias de incrementais)", 30))
    ws = int(num("    Fulls semanais retidos", 4))
    ms = int(num("    Fulls mensais retidos", 12))
    ys = int(num("    Fulls anuais retidos", 5))
    ded = num("    Razão de deduplicação+compressão esperada (x)", 10)
    ovh = num("    Overhead (RAID/FS/reserva) sobre o líquido (%)", 20) / 100
    rows = []
    for a in range(anos + 1):
        f = fe * (1 + cresc) ** a
        logico = f * (ws + ms + ys) + f * chg * dd
        liquido = logico / ded
        bruto = liquido * (1 + ovh)
        rows.append((a, f, logico, liquido, bruto))
    md = ["# Premissas", "| Parâmetro | Valor |", "|---|---|",
          f"| Dados primários (ano 0) | {fe:,.1f} TB |", f"| Crescimento anual | {cresc*100:.1f}% |",
          f"| Taxa de mudança diária | {chg*100:.2f}% |",
          f"| Retenção | {dd} diários, {ws} semanais, {ms} mensais, {ys} anuais |",
          f"| Deduplicação+compressão | {ded:.1f}:1 |", f"| Overhead | {ovh*100:.0f}% |",
          "# Memória de cálculo",
          "- Dados primários no ano N = D0 × (1 + crescimento)^N",
          "- Volume lógico retido = Dados × (fulls semanais + mensais + anuais) + Dados × mudança diária × dias",
          "- Capacidade líquida = Volume lógico ÷ razão de deduplicação",
          "- Capacidade bruta = Capacidade líquida × (1 + overhead)",
          "# Resultado por ano", "| Ano | Dados primários (TB) | Volume lógico (TB) | Líquida (TB) | Bruta (TB) |",
          "|---|---|---|---|---|"]
    md += [f"| {a} | {f:,.1f} | {l:,.1f} | {q:,.1f} | {b:,.1f} |" for a, f, l, q, b in rows]
    md.append(f"\n**Capacidade líquida recomendada ao final do horizonte: {rows[-1][3]:,.1f} TB "
              f"(bruta {rows[-1][4]:,.1f} TB).** Valores em TB decimais (10^12 bytes).")
    md = "\n".join(md)
    mostrar(md, 60)
    salvar(md, "DIMENSIONAMENTO", "SmartSizing — Dimensionamento de Capacidade")


def m_calc(c):
    dados_sessao()
    n = int(num("  SmartCalc — quantas soluções comparar (1–3)", 2))
    sols = []
    for i in range(max(1, min(3, n))):
        print(f"  Solução {i+1}:")
        sols.append(dict(nome=perguntar("    Nome", f"Solução {i+1}"),
                         capex=num("    Investimento inicial CAPEX (R$)", 0),
                         opex=num("    Custo anual OPEX — suporte/energia/espaço/pessoal (R$)", 0),
                         econ=num("    Economia/benefício anual estimado (R$)", 0)))
    md = ["# Comparativo TCO / ROI / Payback",
          "| Solução | CAPEX | OPEX/ano | TCO 3 anos | TCO 5 anos | ROI 3 anos | ROI 5 anos | Payback |",
          "|---|---|---|---|---|---|---|---|"]
    for s in sols:
        t3, t5 = s["capex"] + 3 * s["opex"], s["capex"] + 5 * s["opex"]
        roi = lambda t, a: f"{((s['econ'] * a - t) / t * 100):.1f}%" if t else "n/a"
        liq = s["econ"] - s["opex"]
        pb = f"{s['capex'] / liq * 12:.1f} meses" if liq > 0 else "não se paga"
        md.append(f"| {s['nome']} | R$ {s['capex']:,.2f} | R$ {s['opex']:,.2f} | R$ {t3:,.2f} | R$ {t5:,.2f} "
                  f"| {roi(t3, 3)} | {roi(t5, 5)} | {pb} |")
    md += ["# Memória de cálculo", "- TCO(N) = CAPEX + OPEX anual × N",
           "- ROI(N) = (Benefício anual × N − TCO(N)) ÷ TCO(N)",
           "- Payback = CAPEX ÷ (Benefício anual − OPEX anual)"]
    md = "\n".join(md)
    mostrar(md)
    salvar(md, "TCO_ROI", "SmartCalc — TCO, ROI e Payback")


CAMPOS_INTEL = ["data", "orgao", "pregao", "objeto", "vencedor", "fabricante", "produto", "valor_estimado",
                "valor_vencedor", "nosso_valor", "resultado", "observacoes"]


def m_intel(c):
    print("  SmartIntel — 1) Registrar pregão  2) Listar base  3) Importar CSV")
    op = perguntar("  Opção", "1")
    novo = not INTEL_CSV.exists()
    if op == "1":
        reg = {k: perguntar(f"    {k}") for k in CAMPOS_INTEL}
        with INTEL_CSV.open("a", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=CAMPOS_INTEL)
            if novo:
                w.writeheader()
            w.writerow(reg)
        print(f"  ✔ Registrado em {INTEL_CSV}")
    elif op == "2":
        print(INTEL_CSV.read_text(encoding="utf-8") if INTEL_CSV.exists() else "  (base vazia)")
    elif op == "3":
        p = caminho("CSV com colunas " + ",".join(CAMPOS_INTEL))
        with p.open(encoding="utf-8-sig") as src, INTEL_CSV.open("a", newline="", encoding="utf-8") as dst:
            w = csv.DictWriter(dst, fieldnames=CAMPOS_INTEL, extrasaction="ignore")
            if novo:
                w.writeheader()
            k = 0
            for r in csv.DictReader(src):
                w.writerow(r)
                k += 1
        print(f"  ✔ {k} registros importados.")


# ───────────────────────────── SmartSearch (PNCP) ─────────────────────────────
PNCP_SEARCH = "https://pncp.gov.br/api/search/"
PNCP_API = "https://pncp.gov.br/api/pncp/v1"
PRIORIDADE = re.compile(r"tribunal|minist|agência|agencia|banco|caixa|petrobras|empresa|instituto|universidade|"
                        r"autarquia|conselho|serpro|dataprev|eletro|correios|ibge|receita|polícia|policia|exército|"
                        r"marinha|aeronáutica|justiça|justica|procuradoria|ministério público", re.I)


def _req():
    try:
        import requests
        return requests
    except ImportError:
        sys.exit("Instale as dependências: pip install -r requirements.txt")


def pncp_busca(termo, status, paginas=2):
    rq = _req()
    itens = []
    for pg in range(1, paginas + 1):
        try:
            r = rq.get(PNCP_SEARCH, params={"q": termo, "tipos_documento": "edital", "ordenacao": "-data",
                                            "pagina": pg, "tam_pagina": 50, "status": status},
                       timeout=40, headers={"User-Agent": "SmartBid/1.0"})
            r.raise_for_status()
            lote = r.json().get("items", [])
        except Exception as e:
            print(f"  ! PNCP ({termo}): {e}")
            break
        itens += lote
        if len(lote) < 50:
            break
    return itens


def link_pncp(it):
    u = it.get("item_url") or ""
    m = re.search(r"/compras/(\d+)/(\d+)/(\d+)", u)
    return (f"https://pncp.gov.br/app/editais/{m.group(1)}/{m.group(2)}/{m.group(3)}", m.groups()) if m else ("", None)


def m_search(c):
    print("  SmartSearch — busca no PNCP (Portal Nacional de Contratações Públicas)")
    kws = perguntar("  Palavras-chave (separadas por ;)", "; ".join(c["keywords"]))
    kws = [k.strip() for k in kws.split(";") if k.strip()]
    ufs = [u.strip().upper() for u in perguntar("  UFs (ex.: DF;GO — Enter = todas)",
                                                ";".join(c["ufs"])).split(";") if u.strip()]
    vmin = num("  Valor estimado mínimo (R$)", c["valor_min"])
    orgao_f = perguntar("  Filtro por órgão (contém — Enter = todos)").lower()
    mod_f = perguntar("  Filtro por modalidade (ex.: pregão — Enter = todas)").lower()
    st = {"1": "recebendo_proposta", "2": "encerradas", "3": "todos"}.get(perguntar(
        "  Status: 1) recebendo propostas 2) encerradas 3) todos", "1"), "recebendo_proposta")
    hist = json.loads(HIST_SEARCH.read_text(encoding="utf-8")) if HIST_SEARCH.exists() else {}
    vistos, res = set(), []
    for k in kws:
        print(f"  · buscando '{k}' ...")
        for it in pncp_busca(k, st):
            chave = it.get("numero_controle_pncp") or it.get("item_url") or it.get("id")
            if chave in vistos:
                continue
            vistos.add(chave)
            uf = (it.get("uf") or "").upper()
            valor = float(it.get("valor_global") or 0)
            orgao = it.get("orgao_nome") or ""
            mod = (it.get("modalidade_licitacao_nome") or "").lower()
            if ufs and uf not in ufs:
                continue
            if vmin and valor and valor < vmin:
                continue
            if orgao_f and orgao_f not in orgao.lower():
                continue
            if mod_f and mod_f not in mod:
                continue
            esfera = (it.get("esfera_nome") or "").lower()
            prio = (2 if "federal" in esfera else 0) + (1 if PRIORIDADE.search(orgao) else 0)
            url, ids = link_pncp(it)
            res.append(dict(chave=chave, prio=prio, uf=uf, valor=valor, orgao=orgao,
                            municipio=it.get("municipio_nome") or "", modalidade=it.get("modalidade_licitacao_nome") or "",
                            titulo=it.get("title") or "", objeto=(it.get("description") or "").replace("\n", " "),
                            fim=(it.get("data_fim_vigencia") or "")[:10], url=url, ids=ids, termo=k,
                            novo=chave not in hist))
    res.sort(key=lambda r: (-r["prio"], -r["valor"]))
    if not res:
        print("  Nenhuma oportunidade encontrada com esses filtros.")
        return
    linha()
    for i, r in enumerate(res, 1):
        flag = "🆕" if r["novo"] else "  "
        print(f"{i:>3} {flag} {'★' * r['prio']:<3} [{r['uf']}] {r['orgao'][:55]} | R$ {r['valor']:,.2f} | "
              f"prop. até {r['fim']}\n        {r['objeto'][:150]}\n        {r['url']}")
    linha()
    for r in res:
        hist.setdefault(r["chave"], {"capturado": dt.datetime.now().isoformat(timespec="minutes"),
                                     "status": "novo", "orgao": r["orgao"], "url": r["url"]})
    HIST_SEARCH.write_text(json.dumps(hist, indent=1, ensure_ascii=False), encoding="utf-8")
    arq = OUT / f"SMARTSEARCH_{dt.datetime.now():%d%m%y_%H%M}.csv"
    with arq.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Prioridade", "UF", "Município", "Órgão", "Modalidade", "Objeto", "Valor estimado",
                    "Propostas até", "Link PNCP", "Termo", "Novo"])
        for r in res:
            w.writerow([r["prio"], r["uf"], r["municipio"], r["orgao"], r["modalidade"], r["objeto"],
                        f"{r['valor']:.2f}".replace(".", ","), r["fim"], r["url"], r["termo"],
                        "sim" if r["novo"] else "não"])
    print(f"  ✔ {len(res)} oportunidades ({sum(r['novo'] for r in res)} novas). Planilha: {arq}")
    print("  ★★ = federal  ★ = estatal/autarquia/grande órgão  🆕 = não capturado antes")
    sel = perguntar("  Nº para baixar o edital e rodar o SmartBid completo (Enter = voltar)")
    if sel.isdigit() and 1 <= int(sel) <= len(res):
        r = res[int(sel) - 1]
        edital = baixar_edital(r)
        if edital:
            hist[r["chave"]]["status"] = "em análise"
            HIST_SEARCH.write_text(json.dumps(hist, indent=1, ensure_ascii=False), encoding="utf-8")
            SESSAO["cliente"] = None
            SESSAO["pregao"] = None
            fluxo(c, edital)


def baixar_edital(r):
    if not r["ids"]:
        print("  ! Não foi possível identificar a compra no PNCP.")
        return None
    rq = _req()
    cnpj, ano, seq = r["ids"]
    try:
        arqs = rq.get(f"{PNCP_API}/orgaos/{cnpj}/compras/{ano}/{seq}/arquivos", timeout=40).json()
    except Exception as e:
        print(f"  ! Falha ao listar arquivos: {e}")
        return None
    if not arqs:
        print("  ! Nenhum arquivo publicado para esta compra.")
        return None
    for i, a in enumerate(arqs, 1):
        print(f"   {i}) {a.get('tipoDocumentoNome', '')} — {a.get('titulo', '')}")
    s = perguntar("  Arquivo a baixar", "1")
    a = arqs[int(s) - 1] if s.isdigit() and 1 <= int(s) <= len(arqs) else arqs[0]
    url = a.get("url") or a.get("uri")
    pasta = OUT / "editais"
    pasta.mkdir(exist_ok=True)
    resp = rq.get(url, timeout=120)
    nome = re.findall(r'filename="?([^";]+)', resp.headers.get("content-disposition", ""))
    nome = nome[0] if nome else f"{cnpj}_{ano}_{seq}.pdf"
    destino = pasta / f"{cnpj}_{ano}_{seq}_{Path(nome).name}"
    destino.write_bytes(resp.content)
    if destino.suffix.lower() not in (".pdf", ".zip", ".docx"):
        cab = resp.content[:4]
        destino = destino.rename(destino.with_suffix(".zip" if cab.startswith(b"PK") else ".pdf"))
    print(f"  ✔ Baixado: {destino}")
    return destino


# ───────────────────────────── fluxo SmartBid completo ─────────────────────────────
def fluxo(c, edital=None):
    print("\n  SmartBid — fluxo integrado: Insight → Analyse → Compliance → Position (→ Proposal)")
    dados_sessao()
    edital = edital or caminho("Arquivo do EDITAL/TR (PDF, DOCX, ZIP)")
    SESSAO["contexto"].clear()
    gerados = []
    for nome, fn in (("SmartInsight", m_insight), ("SmartAnalyse", m_analyse),
                     ("SmartCompliance", m_compliance), ("SmartPosition", m_position)):
        print(f"\n▶ {nome}")
        arq, md = fn(c, edital)
        gerados.append(arq)
        if nome == "SmartAnalyse" and re.search(r"N[ÃA]O PARTICIPAR", md.split("# 7")[-1], re.I) \
                and not re.search(r"PARTICIPAR COM RESSALVAS", md.split("# 7")[-1], re.I):
            if perguntar("  Recomendação: NÃO PARTICIPAR. Continuar mesmo assim? (s/n)", "n").lower() != "s":
                break
    if perguntar("\n  Gerar a proposta (SmartProposal) agora? (s/n)", "n").lower() == "s":
        arq, _ = m_proposal(c, edital)
        gerados.append(arq)
    print("\n  ✔ Fluxo concluído. Relatórios:")
    for g in gerados:
        print(f"    - {g}")


# ───────────────────────────── menus ─────────────────────────────
def cab():
    linha("═")
    print("  SmartBid v1.0 — Inteligência em Licitações   |   R2 Tecnologia")
    c = cfg()
    print(f"  Produto: {c['produto']}   |   Modelo: {c['modelo']}")
    if SESSAO["cliente"]:
        print(f"  Sessão: {SESSAO['cliente']} — {SESSAO['pregao']}")
    linha("═")


def submenu(c):
    ops = {"1": ("SmartInsight    — síntese executiva do edital", m_insight),
           "2": ("SmartAnalyse    — análise estratégica e score", m_analyse),
           "3": ("SmartCompliance — matriz de conformidade e questionamentos", m_compliance),
           "4": ("SmartPosition   — inteligência competitiva", m_position),
           "5": ("SmartProposal   — proposta técnica e comercial", m_proposal),
           "6": ("SmartReview     — análise do concorrente e recurso", m_review),
           "7": ("SmartDefense    — contrarrazões ao recurso", m_defense),
           "8": ("SmartSizing     — dimensionamento de capacidade", m_sizing),
           "9": ("SmartCalc       — TCO / ROI / payback", m_calc),
           "10": ("SmartIntel      — base de inteligência competitiva", m_intel)}
    while True:
        print("\n  MÓDULOS")
        for k, (t, _) in ops.items():
            print(f"  {k:>2} - {t}")
        print("   0 - voltar")
        op = perguntar("  Opção")
        if op == "0" or not op:
            return
        if op in ops:
            executar(ops[op][1], c)


def executar(fn, c):
    try:
        fn(c)
    except KeyboardInterrupt:
        print("\n  (cancelado)")
    except SystemExit as e:
        print(f"\n  ! {e}")
    except Exception as e:
        print(f"\n  ! Erro: {type(e).__name__}: {e}")


def main():
    while True:
        cab()
        print("  0 - SmartSearch   — localizar oportunidades no PNCP")
        print("  1 - SmartBid      — fluxo completo a partir do edital")
        print("  2 - Módulos       — Insight, Analyse, Compliance, Position, Proposal,")
        print("                      Review, Defense, Sizing, Calc, Intel")
        print("  3 - Nova sessão   — trocar cliente/pregão")
        print("  9 - Sair")
        op = perguntar("\n  Opção")
        c = cfg()
        if op == "0":
            executar(m_search, c)
        elif op == "1":
            executar(fluxo, c)
        elif op == "2":
            submenu(c)
        elif op == "3":
            SESSAO.update(cliente=None, pregao=None, contexto={})
        elif op == "9":
            break


if __name__ == "__main__":
    main()
